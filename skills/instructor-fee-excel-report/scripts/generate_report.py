#!/usr/bin/env python3
"""
강사료 정산 엑셀 생성기.

계산식 (고정):
    총강사료 = 수업시간 * 시급 * 1.5
    원천징수 = 총강사료 * 3.3%  (원 단위 절사)
    실지급액 = 총강사료 - 원천징수

입력 CSV 컬럼: 강사명, 과목(선택), 수업시간, 시급, 비고(선택)
"""

import argparse
import csv
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RATE_MULTIPLIER = 1.5
WITHHOLDING_RATE = 0.033

HEADERS = ["강사명", "과목", "수업시간", "시급", "총강사료", "원천징수(3.3%)", "실지급액", "비고"]


def load_rows(input_path):
    with open(input_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        return list(reader)


def compute_row(raw):
    hours = float(raw["수업시간"])
    rate = float(raw["시급"])
    total_fee = hours * rate * RATE_MULTIPLIER
    withholding = math.floor(total_fee * WITHHOLDING_RATE)
    net_pay = total_fee - withholding
    return {
        "강사명": raw.get("강사명", ""),
        "과목": raw.get("과목", ""),
        "수업시간": hours,
        "시급": rate,
        "총강사료": total_fee,
        "원천징수(3.3%)": withholding,
        "실지급액": net_pay,
        "비고": raw.get("비고", ""),
    }


def build_workbook(rows, month_label):
    wb = Workbook()
    ws = wb.active
    ws.title = "강사료 정산"

    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = f"{month_label} 강사료 정산표" if month_label else "강사료 정산표"
    title_cell.font = Font(size=14, bold=True)
    title_cell.alignment = Alignment(horizontal="center")

    header_row = 3
    header_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    currency_cols = {"시급", "총강사료", "원천징수(3.3%)", "실지급액"}
    total_fee_sum = 0
    withholding_sum = 0
    net_pay_sum = 0

    r = header_row + 1
    for row in rows:
        computed = compute_row(row)
        for col, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=r, column=col, value=computed[header])
            if header in currency_cols:
                cell.number_format = "#,##0원"
        total_fee_sum += computed["총강사료"]
        withholding_sum += computed["원천징수(3.3%)"]
        net_pay_sum += computed["실지급액"]
        r += 1

    total_row = r
    ws.cell(row=total_row, column=1, value="합계").font = Font(bold=True)
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    ws.cell(row=total_row, column=5, value=total_fee_sum).number_format = "#,##0원"
    ws.cell(row=total_row, column=6, value=withholding_sum).number_format = "#,##0원"
    ws.cell(row=total_row, column=7, value=net_pay_sum).number_format = "#,##0원"
    for col in range(1, len(HEADERS) + 1):
        ws.cell(row=total_row, column=col).font = Font(bold=True)

    widths = [12, 14, 10, 10, 12, 14, 12, 20]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    return wb


def main():
    parser = argparse.ArgumentParser(description="강사료 정산 엑셀 생성")
    parser.add_argument("--input", required=True, help="입력 CSV 파일 경로")
    parser.add_argument("--output", required=True, help="출력 xlsx 파일 경로")
    parser.add_argument("--month", default="", help='표 제목에 넣을 월 표기, 예: "2026년 9월"')
    args = parser.parse_args()

    rows = load_rows(args.input)
    wb = build_workbook(rows, args.month)
    wb.save(args.output)
    print(f"저장 완료: {args.output} ({len(rows)}명)")


if __name__ == "__main__":
    main()
